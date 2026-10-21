#!/usr/bin/env python3
"""
Generador nativo y 100% compatible de usuarios.xlsx para la Plataforma Forta.
Usa openpyxl para garantizar compatibilidad total con Microsoft Excel en Windows/Mac,
Excel Online, Google Sheets y LibreOffice.
"""

import os
import sys
import json
import random
import hashlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
JSON_PATH = os.path.join(BASE_DIR, 'data', 'usuarios.json')
PW_PATH = os.path.join(BASE_DIR, 'tools', '_coord-pw.json')
EXCEL_PATH = os.path.join(BASE_DIR, 'usuarios.xlsx')

PALABRAS = ['casa', 'sol', 'mar', 'luna', 'rio', 'vida', 'flor', 'tono',
            'piedra', 'viento', 'hielo', 'arena', 'palma', 'osbo', 'nieve', 'grano']

def genera_password():
    a = random.choice(PALABRAS)
    b = random.choice(PALABRAS)
    while b == a:
        b = random.choice(PALABRAS)
    num = f"{random.randint(0, 999):03d}"
    # La contraseña empieza en mayúscula
    return f"{a.capitalize()}.{b}.{num}"

def sha256(text):
    return hashlib.sha256(text.encode('utf-8')).hexdigest()

def main():
    regenerate = '--regenerate' in sys.argv

    with open(JSON_PATH, 'r', encoding='utf-8') as f:
        uj = json.load(f)

    pwd_claro = {}
    if os.path.exists(PW_PATH) and not regenerate:
        try:
            with open(PW_PATH, 'r', encoding='utf-8') as f:
                pwd_claro = json.load(f)
        except Exception:
            pwd_claro = {}

    # Asegurar que todas las contraseñas en memoria empiecen en mayúscula
    for usr, p in list(pwd_claro.items()):
        if p and p[0].islower():
            pwd_claro[usr] = p[0].upper() + p[1:]

    changed = False
    for u in uj.get('usuarios', []):
        usr = u['usuario']
        if usr not in pwd_claro or regenerate:
            p = genera_password()
            pwd_claro[usr] = p
            u['password_sha256'] = sha256(p)
            changed = True
        else:
            # Asegurar que el hash en JSON coincida con la contraseña guardada (que empieza en mayúscula)
            p = pwd_claro[usr]
            expected_hash = sha256(p)
            if u.get('password_sha256') != expected_hash:
                u['password_sha256'] = expected_hash
                changed = True

        # Permisos adicionales de municipios para Carlos Mejía:
        # Alpatláhuac (9), Calcahualco (30), Coscomatepec (49), Chocamán (64), Tomatlán (185)
        if usr.lower() == 'carlos.mejia':
            nuevos_extra = [9, 30, 49, 64, 185]
            if u.get('municipios_adicionales') != nuevos_extra:
                u['municipios_adicionales'] = nuevos_extra
                changed = True

    if changed:
        with open(JSON_PATH, 'w', encoding='utf-8') as f:
            json.dump(uj, f, indent=2, ensure_ascii=False)
            f.write('\n')

    # Guardar contraseñas en claro persistentes
    with open(PW_PATH, 'w', encoding='utf-8') as f:
        json.dump(pwd_claro, f, indent=2, ensure_ascii=False)
        f.write('\n')

    # Crear libro de Excel nativo
    wb = Workbook()
    ws = wb.active
    ws.title = "Usuarios y Contraseñas"
    ws.views.sheetView[0].showGridLines = True

    # Estilos institucionales
    azul_header = "003F87"  # Azul marino institucional Forta
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color=azul_header, end_color=azul_header, fill_type="solid")
    align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

    font_data = Font(name="Segoe UI", size=10.5)
    font_bold = Font(name="Segoe UI", size=10.5, bold=True)
    font_pwd = Font(name="Consolas", size=11, bold=True, color="003F87")

    thin_border_side = Side(border_style="thin", color="D0D7DE")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    fill_even = PatternFill(start_color="F9FAFC", end_color="F9FAFC", fill_type="solid")
    fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    headers = [
        "Usuario",
        "Nombre Completo",
        "Rol",
        "Estado",
        "Email Sugerido",
        "Contraseña (Inicia con Mayúscula + 3 Dígitos)"
    ]
    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_header
        cell.border = border_cell

    for row_idx, u in enumerate(uj.get('usuarios', []), start=2):
        usr = u['usuario']
        nombre = u.get('nombre', usr)
        rol = (u.get('rol') or 'admin').capitalize()
        estado = "Inactivo" if u.get('activo') is False else "Activo"
        email = f"{usr.lower()}@platadeforta.mx"
        pwd = pwd_claro.get(usr, '')

        row_data = [usr, nombre, rol, estado, email, pwd]
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 22

        fill_row = fill_even if row_idx % 2 == 0 else fill_odd

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border_cell
            cell.fill = fill_row

            if col_idx == 1:
                cell.font = font_bold
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx in (3, 4):
                cell.font = font_data
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 6:
                cell.font = font_pwd
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.font = font_data
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Autoajuste de ancho de columnas con holgura
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Activar autofiltro
    ws.auto_filter.ref = ws.dimensions

    wb.save(EXCEL_PATH)
    print(f"Excel nativo generado con éxito en: {EXCEL_PATH}")
    print(f"Total de usuarios incluidos: {len(uj.get('usuarios', []))}")

if __name__ == '__main__':
    main()
